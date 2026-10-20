import io
import json
import asyncio
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, status, Query, Request
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, desc
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.db.session import get_db
from app.models.entities import Tenant, Campaign, Candidate, VoterIdentity, VoteLog
from app.schemas.log import VoteLogOut
from app.core.deps import get_current_tenant
from app.services.broadcaster import broadcaster

router = APIRouter(prefix="/analytics", tags=["analytics"])

@router.get("/campaigns/{campaign_id}/stream")
async def campaign_event_stream(
    campaign_id: int,
    request: Request,
    db: Session = Depends(get_db)
):
    campaign = db.query(Campaign).filter(Campaign.id == campaign_id).first()
    if not campaign:
        raise HTTPException(status_code=404, detail="活动未找到")

    queue = broadcaster.subscribe(campaign.id)

    async def event_generator():
        try:
            # Send initial state snapshot immediately
            init_tallies = {c.id: c.vote_count for c in campaign.candidates}
            initial_data = {
                "event": "initial_snapshot",
                "campaign_id": campaign.id,
                "tallies": init_tallies,
                "total_votes": sum(init_tallies.values()),
                "timestamp": datetime.now().isoformat()
            }
            yield f"data: {json.dumps(initial_data, ensure_ascii=False)}\n\n"

            while True:
                # Check client disconnect
                if await request.is_disconnected():
                    break

                try:
                    # Wait up to 15 seconds for a vote event or send heartbeat
                    message = await asyncio.wait_for(queue.get(), timeout=15.0)
                    yield f"data: {json.dumps(message, ensure_ascii=False)}\n\n"
                except asyncio.TimeoutError:
                    # Heartbeat comment to keep SSE connection alive
                    yield ": ping\n\n"
        finally:
            broadcaster.unsubscribe(campaign.id, queue)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )

@router.get("/campaigns/{campaign_id}/stats")
def get_campaign_stats(
    campaign_id: int,
    current_tenant: Tenant = Depends(get_current_tenant),
    db: Session = Depends(get_db)
):
    campaign = db.query(Campaign).filter(
        Campaign.id == campaign_id,
        Campaign.tenant_id == current_tenant.id
    ).first()
    if not campaign:
        raise HTTPException(status_code=404, detail="活动未找到")

    total_valid_votes = db.query(func.count(VoteLog.id)).filter(
        VoteLog.campaign_id == campaign.id,
        VoteLog.is_valid == 1
    ).scalar() or 0

    total_participants = db.query(func.count(VoterIdentity.id)).filter(
        VoterIdentity.campaign_id == campaign.id
    ).scalar() or 0

    unique_ips = db.query(func.count(func.distinct(VoteLog.client_ip))).filter(
        VoteLog.campaign_id == campaign.id
    ).scalar() or 0

    candidates_ranking = db.query(Candidate).filter(
        Candidate.campaign_id == campaign.id
    ).order_by(Candidate.vote_count.desc()).all()

    return {
        "campaign_id": campaign.id,
        "title": campaign.title,
        "total_votes": total_valid_votes,
        "total_participants": total_participants,
        "unique_ips": unique_ips,
        "rankings": [
            {
                "id": c.id,
                "title": c.title,
                "subtitle": c.subtitle,
                "vote_count": c.vote_count,
                "cover_url": c.cover_url
            }
            for c in candidates_ranking
        ]
    }

@router.get("/campaigns/{campaign_id}/logs")
def list_vote_logs(
    campaign_id: int,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    is_valid: Optional[int] = None,
    search: Optional[str] = None,
    current_tenant: Tenant = Depends(get_current_tenant),
    db: Session = Depends(get_db)
):
    campaign = db.query(Campaign).filter(
        Campaign.id == campaign_id,
        Campaign.tenant_id == current_tenant.id
    ).first()
    if not campaign:
        raise HTTPException(status_code=404, detail="活动未找到")

    query = db.query(VoteLog).filter(VoteLog.campaign_id == campaign.id)

    if is_valid is not None:
        query = query.filter(VoteLog.is_valid == is_valid)

    if search:
        search_filter = f"%{search}%"
        query = query.filter(
            (VoteLog.client_ip.like(search_filter)) |
            (VoteLog.visitor_id.like(search_filter))
        )

    total = query.count()
    logs = query.order_by(VoteLog.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()

    items = []
    for l in logs:
        voter_info = json.loads(l.voter_identity.voter_info_json or "{}") if l.voter_identity else {}
        items.append({
            "id": l.id,
            "campaign_id": l.campaign_id,
            "candidate_id": l.candidate_id,
            "candidate_title": l.candidate.title if l.candidate else None,
            "identity_id": l.identity_id,
            "client_ip": l.client_ip,
            "visitor_id": l.visitor_id,
            "user_agent": l.user_agent,
            "is_valid": l.is_valid,
            "voter_info": voter_info,
            "created_at": l.created_at
        })

    return {"total": total, "page": page, "page_size": page_size, "items": items}

@router.post("/campaigns/{campaign_id}/logs/{log_id}/invalidate")
async def invalidate_vote_log(
    campaign_id: int,
    log_id: int,
    current_tenant: Tenant = Depends(get_current_tenant),
    db: Session = Depends(get_db)
):
    campaign = db.query(Campaign).filter(
        Campaign.id == campaign_id,
        Campaign.tenant_id == current_tenant.id
    ).first()
    if not campaign:
        raise HTTPException(status_code=404, detail="活动未找到")

    log = db.query(VoteLog).filter(
        VoteLog.id == log_id,
        VoteLog.campaign_id == campaign.id
    ).first()
    if not log:
        raise HTTPException(status_code=404, detail="投票流水不存在")

    if log.is_valid == 0:
        return {"message": "该票已被标记为作废，无需重复操作"}

    # Invalidate log
    log.is_valid = 0

    # Rollback candidate vote count
    if log.candidate:
        log.candidate.vote_count = max(0, log.candidate.vote_count - 1)

    # Rollback identity total voted count
    if log.voter_identity:
        log.voter_identity.total_voted = max(0, log.voter_identity.total_voted - 1)

    db.commit()

    # Trigger SSE update
    candidate_tallies = {c.id: c.vote_count for c in campaign.candidates}
    await broadcaster.broadcast(
        campaign.id,
        {
            "event": "vote_invalidated",
            "campaign_id": campaign.id,
            "invalidated_log_id": log.id,
            "tallies": candidate_tallies,
            "total_votes": sum(candidate_tallies.values())
        }
    )

    return {"message": "作弊票作废成功，已自动扣减对应作品票数"}

@router.get("/campaigns/{campaign_id}/export")
def export_campaign_excel(
    campaign_id: int,
    current_tenant: Tenant = Depends(get_current_tenant),
    db: Session = Depends(get_db)
):
    campaign = db.query(Campaign).filter(
        Campaign.id == campaign_id,
        Campaign.tenant_id == current_tenant.id
    ).first()
    if not campaign:
        raise HTTPException(status_code=404, detail="活动未找到")

    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="微软雅黑", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB")
    )

    # Sheet 1: 最终排名榜
    ws1 = wb.active
    ws1.title = "作品得票排名榜"
    ws1.append(["排名", "作品编号", "作品名称", "选送单位/副标题", "累计有效得票", "得票率"])

    for col_idx in range(1, 7):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_votes = db.query(func.count(VoteLog.id)).filter(
        VoteLog.campaign_id == campaign.id,
        VoteLog.is_valid == 1
    ).scalar() or 0

    candidates = db.query(Candidate).filter(
        Candidate.campaign_id == campaign.id
    ).order_by(Candidate.vote_count.desc()).all()

    for idx, c in enumerate(candidates, 1):
        rate = f"{(c.vote_count / total_votes * 100):.1f}%" if total_votes > 0 else "0.0%"
        ws1.append([idx, c.id, c.title, c.subtitle or "-", c.vote_count, rate])

    for row in ws1.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Sheet 2: 机构与候选人透视矩阵
    ws2 = wb.create_sheet(title="单位组织透视矩阵")
    logs = db.query(VoteLog).filter(
        VoteLog.campaign_id == campaign.id,
        VoteLog.is_valid == 1
    ).all()

    # Aggregate by unit and candidate
    unit_candidate_counts = {}
    all_units = set()
    for l in logs:
        vinfo = json.loads(l.voter_identity.voter_info_json or "{}") if l.voter_identity else {}
        unit = vinfo.get("unit") or vinfo.get("department") or "未分类/公众"
        all_units.add(unit)
        cand_title = l.candidate.title if l.candidate else f"ID_{l.candidate_id}"
        unit_candidate_counts[(unit, cand_title)] = unit_candidate_counts.get((unit, cand_title), 0) + 1

    cand_titles = [c.title for c in candidates]
    ws2.append(["组织/支行单位"] + cand_titles + ["该单位总投出票数"])
    for col_idx in range(1, len(cand_titles) + 3):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for unit in sorted(all_units):
        row_data = [unit]
        unit_total = 0
        for title in cand_titles:
            cnt = unit_candidate_counts.get((unit, title), 0)
            row_data.append(cnt)
            unit_total += cnt
        row_data.append(unit_total)
        ws2.append(row_data)

    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Sheet 3: 完整明细审计流水
    ws3 = wb.create_sheet(title="投票明细审计流水")
    ws3.append(["流水ID", "投票时间 (UTC)", "所投作品", "投票人姓名", "所属单位", "IP地址", "设备指纹", "票状态"])
    for col_idx in range(1, 9):
        cell = ws3.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    all_logs = db.query(VoteLog).filter(
        VoteLog.campaign_id == campaign.id
    ).order_by(VoteLog.created_at.desc()).all()

    for l in all_logs:
        vinfo = json.loads(l.voter_identity.voter_info_json or "{}") if l.voter_identity else {}
        ws3.append([
            l.id,
            l.created_at.strftime("%Y-%m-%d %H:%M:%S"),
            l.candidate.title if l.candidate else f"ID:{l.candidate_id}",
            vinfo.get("real_name") or "-",
            vinfo.get("unit") or vinfo.get("department") or "-",
            l.client_ip,
            l.visitor_id[:16] + "...",
            "有效" if l.is_valid == 1 else "已作废"
        ])

    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    # Auto-adjust column widths
    for sheet in [ws1, ws2, ws3]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(max_len + 3, 12), 45)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"campaign_{campaign.code}_audit_report_{datetime.now().strftime('%Y%m%d%H%M')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
